"""Карточка организации (реквизиты) в фирменном стиле Creatica — как КП: монохром, тонкие линии, IBM Plex Sans.

Реквизиты не хранятся в публичном репозитории: они лежат в company_cards.json (приватный пакет
PRIVATE_BUNDLE_B64, как principal.py). Если файла нет — берётся company_cards_example.json с заглушками.
"""

from __future__ import annotations

import io
import json
import os

from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader, simpleSplit
from reportlab.pdfgen import canvas

from utils import tag as tagmod
from utils.kp import BRAND, GREY, HAIR, INK, MR, ML, PH, PW, SHOWROOM, SITE

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def load_cards() -> dict[str, dict]:
    """{ключ: {"short": ..., "full": ..., "rows": [[подпись, значение], ...]}} в порядке файла."""
    for name in ("company_cards.json", "company_cards_example.json"):
        path = os.path.join(ROOT, name)
        if os.path.exists(path):
            with open(path, encoding="utf-8") as f:
                return json.load(f)
    return {}


def card_text(card: dict) -> str:
    """Реквизиты текстом — чтобы скопировать и переслать клиенту одним сообщением."""
    lines = [card["full"], ""] + [f"{label}: {value}" for label, value in card["rows"]]
    return "\n".join(lines)


def build_card_pdf(card: dict) -> bytes:
    tagmod._register_fonts()
    reg, bold = tagmod.REG, tagmod.BOLD
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setTitle(f"Карточка предприятия — {card['short']}")
    c.setAuthor("Creatica")

    def Y(y: float) -> float:
        return PH - y

    def text(x, y, s, size=9.5, is_bold=False, color=INK, space=0.0, align="l"):
        font = bold if is_bold else reg
        c.setFillColorRGB(*color)
        if space:
            t = c.beginText(x, Y(y))
            t.setFont(font, size)
            t.setCharSpace(space)
            t.textOut(s)
            t.setCharSpace(0)
            c.drawText(t)
            return
        c.setFont(font, size)
        (c.drawRightString if align == "r" else c.drawString)(x, Y(y), s)

    def label(x, y, s):
        text(x, y, s.upper(), size=6.8, color=GREY, space=1.1)

    def hair(y, color=HAIR, width=0.5):
        c.setStrokeColorRGB(*color)
        c.setLineWidth(width)
        c.line(ML, Y(y), MR, Y(y))

    # шапка — как в КП: логотип слева, сайт справа
    logo = ImageReader(os.path.join(BRAND, "creatica_logo.png"))
    h = 17.0
    w = h * logo.getSize()[0] / logo.getSize()[1]
    c.drawImage(logo, ML, Y(34 + h), w, h, mask="auto")
    text(MR, 45, SITE, size=8.5, color=GREY, align="r")

    y = 106.0
    label(ML, y, "Карточка предприятия")
    y += 14
    hair(y, color=INK, width=0.8)
    y += 40
    text(ML, y, card["short"], size=22, is_bold=True)
    y += 20
    for ln in simpleSplit(card["full"], reg, 10.5, MR - ML):
        text(ML, y, ln, size=10.5, color=GREY)
        y += 14
    y += 18
    hair(y)

    # реквизиты: подпись слева, значение справа, тонкая линия между строками
    col = ML + 170
    for lab, value in card["rows"]:
        lines = []
        for part in str(value).split("\n"):
            lines += simpleSplit(part, reg, 10.5, MR - col) or [""]
        row_h = 16 + 14 * len(lines)
        label(ML, y + 19, lab)
        for i, ln in enumerate(lines):
            text(col, y + 20 + i * 14, ln, size=10.5)
        y += row_h + 4
        hair(y)

    # чёрная плашка внизу — фирменный акцент, как в конце КП
    panel_h = 70.0
    top = PH - 52 - panel_h
    c.setFillColorRGB(*INK)
    c.rect(0, Y(top + panel_h), PW, panel_h, stroke=0, fill=1)
    text(ML, top + 32, "Creatica", size=18, is_bold=True, color=(1, 1, 1))
    text(ML, top + 50, SITE, size=9.5, color=(0.86, 0.85, 0.83))
    text(MR, top + 32, card["short"], size=9.5, color=(0.86, 0.85, 0.83), align="r")

    # подвал — как в КП
    hair(PH - 38)
    text(ML, PH - 26, SHOWROOM, size=7, color=GREY)
    c.showPage()
    c.save()
    return buf.getvalue()


def build_card_docx(card: dict) -> bytes:
    """Та же карточка в Word: логотип, крупное название, реквизиты с тонкими линиями, чёрная плашка."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    FONT = "Arial"  # есть у любого получателя; IBM Plex в Word у клиента может не оказаться
    ink, grey, hairc = RGBColor(0x12, 0x12, 0x12), RGBColor(0x87, 0x82, 0x7A), "D9D4CC"
    content_w = MR - ML

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(PW), Pt(PH)
    sec.left_margin = sec.right_margin = Pt(ML)
    sec.top_margin, sec.bottom_margin = Pt(34), Pt(40)
    sec.footer_distance = Pt(20)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_before = st.paragraph_format.space_after = Pt(0)

    def run(p, text, size, bold=False, color=ink, spacing=0.0):
        r = p.add_run(text)
        r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
        if spacing:
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(int(spacing * 20)))
            r._r.get_or_add_rPr().append(sp)
        return r

    def para_border(p, color, size_eighths):
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(size_eighths))
        b.set(qn("w:space"), "6"); b.set(qn("w:color"), color)
        bdr.append(b)
        pPr.append(bdr)

    def cell_style(cell, bottom=None, fill=None, pad=(0, 0, 0, 0)):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{side}")
            if side == "bottom" and bottom:
                e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), bottom)
            else:
                e.set(qn("w:val"), "nil")
            borders.append(e)
        tcPr.append(borders)
        if fill:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
            tcPr.append(shd)
        mar = OxmlElement("w:tcMar")
        for side, val in zip(("top", "left", "bottom", "right"), pad):
            m = OxmlElement(f"w:{side}")
            m.set(qn("w:w"), str(int(val * 20))); m.set(qn("w:type"), "dxa")
            mar.append(m)
        tcPr.append(mar)

    def table(cols: list[float]):
        t = doc.add_table(rows=0, cols=len(cols))
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False
        for i, w in enumerate(cols):
            t.columns[i].width = Pt(w)
        return t

    def set_widths(row, cols):
        for cell, w in zip(row.cells, cols):
            cell.width = Pt(w)

    # шапка: логотип слева, сайт справа
    cols = [content_w / 2, content_w / 2]
    head = table(cols)
    row = head.add_row()
    set_widths(row, cols)
    for cell in row.cells:
        cell_style(cell)
    row.cells[0].paragraphs[0].add_run().add_picture(os.path.join(BRAND, "creatica_logo.png"), height=Pt(17))
    p = row.cells[1].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(p, SITE, 8.5, color=grey)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(40)
    run(p, "КАРТОЧКА ПРЕДПРИЯТИЯ", 7, color=grey, spacing=1.1)
    para_border(p, "121212", 6)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    run(p, card["short"], 22, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(18)
    run(p, card["full"], 10.5, color=grey)

    # реквизиты
    cols = [170.0, content_w - 170.0]
    reqs = table(cols)
    for i, (lab, value) in enumerate(card["rows"]):
        row = reqs.add_row()
        set_widths(row, cols)
        for cell in row.cells:
            cell_style(cell, bottom=hairc, pad=(9, 0, 9, 0))
        run(row.cells[0].paragraphs[0], lab.upper(), 7, color=grey, spacing=1.1)
        run(row.cells[1].paragraphs[0], str(value), 10.5)
    top_border = OxmlElement("w:top")  # линия над первой строкой — как в PDF
    top_border.set(qn("w:val"), "single"); top_border.set(qn("w:sz"), "4"); top_border.set(qn("w:color"), hairc)
    for cell in reqs.rows[0].cells:
        cell._tc.tcPr.find(qn("w:tcBorders")).replace(cell._tc.tcPr.find(qn("w:tcBorders")).find(qn("w:top")),
                                                      copy_el(top_border))

    # чёрная плашка
    doc.add_paragraph().paragraph_format.space_after = Pt(30)
    from docx.enum.text import WD_TAB_ALIGNMENT

    band = table([content_w])
    # Word сдвигает таблицу влево на внутренний отступ ячейки — возвращаем плашку ровно к полю страницы
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), str(18 * 20)); ind.set(qn("w:type"), "dxa")
    look = band._tbl.tblPr.find(qn("w:tblLook"))
    (look.addprevious if look is not None else band._tbl.tblPr.append)(ind)
    row = band.add_row()
    set_widths(row, [content_w])
    white, light = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xDB, 0xD9, 0xD4)
    cell_style(row.cells[0], fill="121212", pad=(16, 18, 16, 18))
    p = row.cells[0].paragraphs[0]
    p.paragraph_format.tab_stops.add_tab_stop(Pt(content_w - 36), WD_TAB_ALIGNMENT.RIGHT)
    run(p, "Creatica", 18, bold=True, color=white)
    run(p, "\t" + card["short"], 9.5, color=light)
    run(row.cells[0].add_paragraph(), SITE, 9.5, color=light)

    # подвал
    fp = sec.footer.paragraphs[0]
    run(fp, SHOWROOM, 7, color=grey)
    para_border_top = OxmlElement("w:pBdr")
    t = OxmlElement("w:top")
    t.set(qn("w:val"), "single"); t.set(qn("w:sz"), "4"); t.set(qn("w:space"), "6"); t.set(qn("w:color"), hairc)
    para_border_top.append(t)
    fp._p.get_or_add_pPr().append(para_border_top)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def copy_el(el):
    import copy

    return copy.deepcopy(el)


def build_card_preview(card: dict, dpi: int = 70) -> bytes:
    import pymupdf

    doc = pymupdf.open(stream=build_card_pdf(card), filetype="pdf")
    return doc[0].get_pixmap(dpi=dpi).tobytes("png")
