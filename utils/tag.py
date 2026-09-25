"""Ценник на мебель, А4: PDF (ReportLab, координаты снятые с шаблона) и Word (python-docx)."""

from __future__ import annotations

import io
import os
from dataclasses import dataclass

from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FONTS = os.path.join(ROOT, "fonts")

SITE = "creatica.shop"
INSTALLMENT_MONTHS = 6
NBSP = " "

# Шаблон снят в пикселях превью 1414x2000 -> пункты А4
K = 595.28 / 1414
W, H = A4

X_LEFT = 95 * K
X_VALUE = 680 * K
X_RIGHT = 1319 * K
X_VLINE = 586 * K
Y_HLINE = 1414

_FONT_OK = False
REG, BOLD = "TagRegular", "TagBold"


def _register_fonts() -> None:
    global REG, BOLD, _FONT_OK
    if _FONT_OK:
        return
    try:
        pdfmetrics.registerFont(TTFont("TagRegular", os.path.join(FONTS, "IBMPlexSans-Regular.ttf")))
        pdfmetrics.registerFont(TTFont("TagBold", os.path.join(FONTS, "IBMPlexSans-Bold.ttf")))
    except Exception:
        REG, BOLD = "Helvetica", "Helvetica-Bold"
    _FONT_OK = True


@dataclass
class Tag:
    title: str
    length: str = ""
    depth: str = ""
    height: str = ""
    frame: str = ""
    filler: str = ""
    article: str = ""
    price: int | None = None
    old_price: int | None = None

    @property
    def installment(self) -> int | None:
        return round(self.price / INSTALLMENT_MONTHS) if self.price else None


def money(value: int | float | None) -> str:
    if value is None:
        return ""
    return f"{int(round(value)):,}".replace(",", NBSP)


def _y(px: float) -> float:
    """Ордината превью (сверху вниз) -> ордината PDF."""
    return H - px * K


def _logo(c: canvas.Canvas, x: float, y_top: float, size: float) -> None:
    """Знак: чёрный квадрат, белый круг, справа щель."""
    y = y_top - size
    c.setFillColorRGB(0, 0, 0)
    c.rect(x, y, size, size, stroke=0, fill=1)
    c.setFillColorRGB(1, 1, 1)
    c.circle(x + 0.501 * size, y + (1 - 0.499) * size, 0.413 * size, stroke=0, fill=1)
    c.rect(x + 0.80 * size, y + (1 - 0.587) * size, 0.21 * size, 0.176 * size, stroke=0, fill=1)
    c.setFillColorRGB(0, 0, 0)


def build_tag_pdf(tag: Tag) -> bytes:
    _register_fonts()
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setTitle(tag.title or "Ценник")
    c.setFillColorRGB(0, 0, 0)

    # сайт
    c.setFont(REG, 13.6)
    c.drawRightString(X_RIGHT, _y(121), SITE)

    # название: подгоняем кегль, затем перенос на 2 строки
    max_w = (1100 - 95) * K
    title = (tag.title or "").strip()
    size = 46.0
    while size > 32 and pdfmetrics.stringWidth(title, REG, size) > max_w:
        size -= 1
    lines = [title]
    if pdfmetrics.stringWidth(title, REG, size) > max_w:
        words, first = title.split(), ""
        for i, word in enumerate(words):
            cand = (first + " " + word).strip()
            if pdfmetrics.stringWidth(cand, REG, size) > max_w and first:
                lines = [first, " ".join(words[i:])]
                break
            first = cand
        size = min(size, 36)
    c.setFont(REG, size)
    y = _y(185)
    for line in lines:
        c.drawString(X_LEFT, y, line)
        y -= size * 1.12

    body = 17.4
    c.setFont(BOLD, 17.4)
    c.drawString(X_LEFT, _y(493), "Габариты")
    c.setFont(REG, body)
    for i, (label, value) in enumerate(
        (("Длина (мм)", tag.length), ("Глубина (мм)", tag.depth), ("Высота (мм)", tag.height))
    ):
        base = _y(595 + i * 51.3)
        c.drawString(X_LEFT, base, label)
        c.drawString(X_VALUE, base, value or "")

    c.setFont(BOLD, 17.4)
    c.drawString(X_LEFT, _y(799), "Характеристики")
    c.setFont(REG, body)
    for i, (label, value) in enumerate((("Каркас", tag.frame), ("Наполнитель", tag.filler))):
        base = _y(900 + i * 51.3)
        c.drawString(X_LEFT, base, label)
        c.drawString(X_VALUE, base, value or "")

    # линии
    c.setStrokeColorRGB(0, 0, 0)
    c.setLineWidth(0.9)
    c.line(0, _y(Y_HLINE), W, _y(Y_HLINE))
    c.line(X_VLINE, _y(Y_HLINE), X_VLINE, 0)

    # низ слева
    c.setFont(REG, 12.6)
    c.drawString(X_LEFT, _y(1535), f"Артикул {tag.article}".rstrip())
    _logo(c, X_LEFT, _y(1812), 93 * K)

    # низ справа
    if tag.installment:
        c.setFont(REG, 13.3)
        c.drawString(X_VALUE, _y(1535), "В рассрочку")
        c.drawString(X_VALUE, _y(1574), f"от {tag.installment} руб./месяц")
    if tag.old_price:
        c.setFont(REG, 18.2)
        text = money(tag.old_price)
        c.drawString(X_VALUE, _y(1765), text)
        tw = pdfmetrics.stringWidth(text, REG, 18.2)
        c.setLineWidth(0.9)
        c.line(X_VALUE - 6, _y(1747), X_VALUE + tw + 6, _y(1747))
    if tag.price:
        c.setFont(REG, 46.5)
        c.drawString(X_VALUE, _y(1905), money(tag.price))

    c.showPage()
    c.save()
    return buf.getvalue()


def build_tag_preview(tag: Tag, dpi: int = 75) -> bytes:
    """Картинка (PNG) ценника — чтобы менеджер видел, как всё расположено."""
    import pymupdf

    doc = pymupdf.open(stream=build_tag_pdf(tag), filetype="pdf")
    return doc[0].get_pixmap(dpi=dpi).tobytes("png")


# ---------------------------------------------------------------- Word


def _logo_png() -> bytes:
    from PIL import Image, ImageDraw

    s = 600
    img = Image.new("RGB", (s * 2, s * 2), "white")
    d = ImageDraw.Draw(img)
    S = s * 2
    d.rectangle([0, 0, S - 1, S - 1], fill="black")
    r = 0.413 * S
    cx, cy = 0.501 * S, 0.499 * S
    d.ellipse([cx - r, cy - r, cx + r, cy + r], fill="white")
    d.rectangle([0.80 * S, 0.411 * S, S, 0.587 * S], fill="white")
    img = img.resize((s, s), Image.LANCZOS)
    out = io.BytesIO()
    img.save(out, "PNG")
    return out.getvalue()


def build_tag_docx(tag: Tag) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    FONT = "Arial"
    BASE = 0.78  # доля высоты строки (точное межстрочное) до базовой линии

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(W), Pt(H)
    sec.left_margin, sec.right_margin = Pt(X_LEFT), Pt(W - X_RIGHT)
    sec.top_margin = sec.bottom_margin = Pt(0)
    sec.header_distance = sec.footer_distance = Pt(0)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)

    def fmt(p, before, line, align=None):
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(max(before, 0)), Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line)
        if align is not None:
            p.alignment = align

    def run(p, text, size, bold=False, strike=False):
        r = p.add_run(text)
        r.font.size, r.bold, r.font.name, r.font.strike = Pt(size), bold, FONT, strike
        return r

    def borders(el, **sides):
        b = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{side}")
            if sides.get(side):
                e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:color"), "000000")
            else:
                e.set(qn("w:val"), "nil")
            b.append(e)
        el.append(b)

    def cell_margins(tcPr, left):
        mar = OxmlElement("w:tcMar")
        for side, val in (("top", 0), ("left", left), ("bottom", 0), ("right", 0)):
            m = OxmlElement(f"w:{side}")
            m.set(qn("w:w"), str(int(val * 20))); m.set(qn("w:type"), "dxa")
            mar.append(m)
        tcPr.append(mar)

    def base_at(px, line):
        """space_before, чтобы базовая линия абзаца оказалась на ординате превью px"""
        return px * K - BASE * line

    # ---- шапка: название слева, сайт справа (таблица без рамок)
    title = (tag.title or "").strip()
    max_w = (1100 - 95) * K
    tsize = 46
    while tsize > 32 and pdfmetrics.stringWidth(title, "Helvetica", tsize) > max_w:
        tsize -= 1  # Arial и Helvetica имеют одинаковые метрики
    two_lines = pdfmetrics.stringWidth(title, "Helvetica", tsize) > max_w
    if two_lines:
        tsize = 34
    tline = tsize * 1.2
    head_h = 205 * K + (tline if two_lines else 0)
    head = doc.add_table(rows=1, cols=2)
    head.autofit = False
    left_w, right_w = 1100 * K - X_LEFT, X_RIGHT - 1100 * K
    for cell, w, lm in ((head.rows[0].cells[0], left_w, 0), (head.rows[0].cells[1], right_w, 0)):
        cell.width = Pt(w)
        cell_margins(cell._tc.get_or_add_tcPr(), lm)
    trPr = head.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight"); h.set(qn("w:val"), str(int(head_h * 20))); h.set(qn("w:hRule"), "exact")
    trPr.append(h)
    p = head.rows[0].cells[0].paragraphs[0]
    fmt(p, base_at(185, tline), tline)
    run(p, title, tsize)
    p = head.rows[0].cells[1].paragraphs[0]
    fmt(p, base_at(121, 18), 18, WD_ALIGN_PARAGRAPH.RIGHT)
    run(p, SITE, 13.5)
    cursor = head_h

    value_tab = Pt(X_VALUE - X_LEFT)

    def line_para(text, baseline_px, size, bold=False, tab_value=None, line=None):
        nonlocal cursor
        line = line or size * 1.3
        p = doc.add_paragraph()
        fmt(p, base_at(baseline_px, line) - cursor, line)
        cursor = max(cursor, 0) + max(base_at(baseline_px, line) - cursor, 0) + line
        if tab_value is not None:
            p.paragraph_format.tab_stops.add_tab_stop(value_tab)
            run(p, f"{text}\t{tab_value or ''}", size)
        else:
            run(p, text, size, bold=bold)

    line_para("Габариты", 493, 18, bold=True)
    for i, (lab, val) in enumerate((("Длина (мм)", tag.length), ("Глубина (мм)", tag.depth), ("Высота (мм)", tag.height))):
        line_para(lab, 595 + i * 51.3, 17, tab_value=val, line=21.6)
    line_para("Характеристики", 799, 18, bold=True)
    for i, (lab, val) in enumerate((("Каркас", tag.frame), ("Наполнитель", tag.filler))):
        line_para(lab, 900 + i * 51.3, 17, tab_value=val, line=21.6)

    # ---- нижний блок: плавающая таблица во всю ширину страницы
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    tbl = table._tbl
    pr = tbl.tblPr
    pos = OxmlElement("w:tblpPr")
    for k, v in (
        ("w:leftFromText", "0"), ("w:rightFromText", "0"), ("w:vertAnchor", "page"),
        ("w:horzAnchor", "page"), ("w:tblpX", str(int(round(X_LEFT * 20)))), ("w:tblpY", str(int(round(Y_HLINE * K * 20)))),
    ):
        pos.set(qn(k), v)
    style_el = pr.find(qn("w:tblStyle"))
    if style_el is not None:
        style_el.addnext(pos)
    else:
        pr.insert(0, pos)
    tblW = pr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        pos.addnext(tblW)
    tblW.set(qn("w:w"), str(int(W * 20))); tblW.set(qn("w:type"), "dxa")
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed")
    look = pr.find(qn("w:tblLook"))
    (look.addprevious(lay) if look is not None else pr.append(lay))

    row_h = H - Y_HLINE * K
    trPr = table.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight"); h.set(qn("w:val"), str(int(row_h * 20))); h.set(qn("w:hRule"), "exact")
    trPr.append(h)

    c1, c2 = table.rows[0].cells
    c1.width, c2.width = Pt(X_VLINE), Pt(W - X_VLINE)
    for cell, lm, right in ((c1, X_LEFT, True), (c2, X_VALUE - X_VLINE, False)):
        tcPr = cell._tc.get_or_add_tcPr()
        borders(tcPr, top=True, right=right)
        cell_margins(tcPr, lm)

    def cell_line(cell, first, text, baseline_px, size, cur, strike=False):
        """абзац в ячейке: базовая линия на baseline_px, cur — текущая ордината (pt от верха линии)"""
        line = size * 1.3
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        before = (baseline_px - Y_HLINE) * K - BASE * line - cur
        fmt(p, before, line)
        if text:
            run(p, text, size, strike=strike)
        return p, max(cur + max(before, 0), 0) + line

    cur = 0.0
    p, cur = cell_line(c1, True, f"Артикул {tag.article}".rstrip(), 1535, 12.5, cur)
    # знак: абзац с картинкой, верх картинки на 1812
    p = c1.add_paragraph()
    pic_top = (1812 - Y_HLINE) * K
    fmt(p, pic_top - cur, 93 * K)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.add_run().add_picture(io.BytesIO(_logo_png()), width=Pt(93 * K), height=Pt(93 * K))

    cur = 0.0
    first = True
    if tag.installment:
        p, cur = cell_line(c2, first, "В рассрочку", 1535, 13, cur); first = False
        p, cur = cell_line(c2, first, f"от {tag.installment} руб./месяц", 1574, 13, cur)
    if tag.old_price:
        p, cur = cell_line(c2, first, money(tag.old_price), 1765, 18, cur, strike=True); first = False
    if tag.price:
        p, cur = cell_line(c2, first, money(tag.price), 1905, 46, cur); first = False

    # Word требует абзац после таблицы — делаем его крошечным
    tail = doc.add_paragraph()
    fmt(tail, 0, 1)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
